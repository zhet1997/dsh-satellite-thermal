"""Conservative, traceable extraction; unstructured sources remain available to DSH.

Numbers are only applied when a column/label identifies the quantity and an
exact unique component ID/name identifies its owner. Free text is never spread
over arbitrary components. An agent may interpret richer prose through PATCH.
"""
from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

from .geometry import safe_zip
from .schema import AXES, issue, origin
from .storage import DomainError

FIELD_ALIASES = {
    "id": ["id", "componentid", "组件id", "部件id"],
    "name": ["name", "component", "componentname", "名称", "组件", "组件名称", "部件", "设备", "设备名称"],
    "material": ["material", "材料", "材质"],
    "conductivity": ["conductivity", "k", "导热率", "热导率", "导热系数"],
    "density": ["density", "密度"], "specificHeat": ["specificheat", "比热", "比热容"],
    "emissivity": ["emissivity", "发射率", "红外发射率"], "absorptivity": ["absorptivity", "吸收率", "太阳吸收率"],
    "power": ["power", "powerw", "功率", "功耗", "发热", "发热量", "热耗"],
    "thickness": ["thickness", "厚度", "板厚"], "kind": ["kind", "类型", "模型类型"],
    "temperatureLimitMin": ["temperaturelimitmin", "最低温度", "温限下限"],
    "temperatureLimitMax": ["temperaturelimitmax", "最高温度", "温限上限"],
}


def normalize_name(text):
    return re.sub(r"[\s_-]+", "", str(text).strip()).casefold()


def field_of(label):
    plain = re.sub(r"[（(\[].*$", "", str(label)).strip()
    return next((key for key, aliases in FIELD_ALIASES.items() if normalize_name(plain) in aliases), None)


def value_with_unit(value, field, header=""):
    if field in {"name", "id", "material", "kind"}:
        return str(value).strip()
    if isinstance(value, (float, int)):
        number, suffix = value, header
    else:
        match = re.fullmatch(r"\s*([-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?)\s*(.*?)\s*", str(value))
        if not match:
            return value
        number, suffix = float(match.group(1)), match.group(2) or header
    if field.startswith("temperatureLimit") and ("°C" in suffix or "℃" in suffix or "摄氏" in suffix):
        number += 273.15
    if field == "power":
        if re.search(r"\bkW\b", suffix):
            number *= 1000
        elif re.search(r"\bmW\b", suffix):
            number /= 1000
    if field == "thickness":
        if re.search(r"\bcm\b", suffix):
            number *= 10
        elif re.search(r"\bm\b", suffix):
            number *= 1000
    if field == "density" and ("g/cm" in suffix or "g·cm" in suffix):
        number *= 1000
    if field == "specificHeat" and "kJ" in suffix:
        number *= 1000
    if field == "conductivity" and "mW" in suffix:
        number /= 1000
    return number


def normalize_row(row, source):
    out = {"_source": source}
    for header, value in row.items():
        field = field_of(header)
        if field and value is not None and str(value).strip() != "":
            out[field] = value_with_unit(value, field, str(header))
    return out if any(k in out for k in ("name", "id")) else None


def rows_from_table(rows, source):
    if not rows:
        return []
    header = [str(v or "") for v in rows[0]]
    records = []
    for idx, values in enumerate(rows[1:], 2):
        row = normalize_row(dict(zip(header, values)), f"{source} 第 {idx} 行")
        if row:
            records.append(row)
    return records


def text_records(text, source, component_names):
    records = []
    aliases = sorted([(alias, field) for field, values in FIELD_ALIASES.items() if field not in {"name", "id"} for alias in values], key=lambda x: -len(x[0]))
    for line_no, line in enumerate(text.splitlines(), 1):
        matches = [name for name in component_names if re.search(r"(?<![\w])" + re.escape(name) + r"(?![\w])", line, re.I)]
        # Chinese names often touch punctuation; the identity must still be exact.
        if len(matches) != 1:
            continue
        row = {"name": matches[0], "_source": f"{source} 第 {line_no} 行"}
        for alias, field in aliases:
            if field in row:
                continue
            pattern = re.escape(alias) + r"\s*[:：=]?\s*([-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?\s*(?:kW|mW|W|K|℃|°C|mm|cm)?)"
            match = re.search(pattern, line, re.I)
            if match:
                row[field] = value_with_unit(match.group(1), field)
        if len(row) > 2:
            records.append(row)
    return records


def extract(path: Path, name, components):
    suffix = path.suffix.lower()
    records, text, cases = [], "", []
    if suffix == ".json":
        try:
            obj = json.loads(path.read_text("utf-8-sig"))
        except (ValueError, UnicodeError) as exc:
            raise DomainError("JSON 格式无效。", "invalid_document") from exc
        rows = obj if isinstance(obj, list) else obj.get("components", [])
        if isinstance(rows, dict):
            rows = [dict(value, name=key) for key, value in rows.items() if isinstance(value, dict)]
        for i, row in enumerate(rows):
            if isinstance(row, dict):
                norm = normalize_row(row, f"{name} components[{i}]")
                if norm:
                    records.append(norm)
        if isinstance(obj, dict) and isinstance(obj.get("cases"), list):
            cases = obj["cases"]
        text = json.dumps(obj, ensure_ascii=False, indent=2)
    elif suffix in {".csv", ".tsv"}:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeError:
            text = raw.decode("gb18030")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        records = [r for i, row in enumerate(csv.DictReader(io.StringIO(text), dialect=dialect), 2) if (r := normalize_row(row, f"{name} 第 {i} 行"))]
        # Explicit six-direction columns are a case-series import, never component power.
        raw_rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
        if raw_rows and all(axis in raw_rows[0] for axis in AXES):
            for i, row in enumerate(raw_rows):
                cases.append({"id": f"import-case-{i}", "name": row.get("name", f"导入工况 {i + 1}"), "kind": "manual",
                              "time": row.get("time", row.get("time_s", i)), "flux": {a: row[a] for a in AXES}, "source": f"{name} 第 {i+2} 行",
                              "valueKind": row.get("valueKind", "incident-total")})
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        # Validate OOXML archive before passing it to a document parser.
        validate_office(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        snippets = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))[:10000]
            records.extend(rows_from_table(rows, f"{name}/{sheet.title}"))
            snippets.extend("\t".join(str(v or "") for v in row) for row in rows)
        workbook.close()
        text = "\n".join(snippets)
    elif suffix == ".docx":
        from docx import Document
        validate_office(path)
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        for i, table in enumerate(doc.tables):
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            records.extend(rows_from_table(rows, f"{name} 表 {i + 1}"))
            text += "\n" + "\n".join("\t".join(row) for row in rows)
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise DomainError("请提供未加密的文字型 PDF。", "encrypted_pdf")
        if len(reader.pages) > 500:
            raise DomainError("文档页数超过 500 页。", "document_too_large", 413)
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
    elif suffix in {".txt", ".md"}:
        text = path.read_text("utf-8-sig")
    else:
        raise DomainError("支持 Excel .xlsx、CSV、Word .docx、文字 PDF、JSON 和文本。", "unsupported_document")
    if len(text) > 5_000_000:
        raise DomainError("提取的文本超过上限，请分批导入。", "document_too_large", 413)
    if suffix in {".txt", ".md", ".docx", ".pdf"}:
        records.extend(text_records(text, name, [c["name"] for c in components]))
    return {"text": text, "records": records, "cases": cases}


def validate_office(path):
    import zipfile
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        if sum(i.file_size for i in infos) > 100 * 1024 * 1024 or len(infos) > 20000:
            raise DomainError("Office 文档解压后过大。", "document_too_large", 413)


def apply_records(project, records):
    applied, unmatched = 0, []
    for row in records:
        identity = row.get("id") or row.get("name")
        matches = [c for c in project["components"] if identity in {c["id"], c.get("sourceId")} or normalize_name(identity) == normalize_name(c["name"])]
        if len(matches) != 1:
            unmatched.append(str(identity))
            continue
        component = matches[0]
        for key, value in row.items():
            if key not in FIELD_ALIASES or key in {"id", "name"}:
                continue
            component[key] = value
            component.setdefault("provenance", {})[key] = origin("provided", row.get("_source", "导入资料"), value)
            applied += 1
    if unmatched:
        project.setdefault("_importIssues", []).append(issue("资料中以下组件未唯一匹配，原文已保留，可通过聊天/参数表匹配：" + "、".join(unmatched[:20]), field="documents"))
    return applied, unmatched
